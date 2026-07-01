"""Word/Excel 文档导入：扫描、解析、写入画布。"""

from __future__ import annotations

import hashlib
import json
import shutil
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from sqlalchemy.orm import Session

from models.canvas_project import CanvasProject, utcnow
from models.excel_import_log import ExcelImportLog
from services.excel_shot_parser import (
    attach_self_check,
    classify_sheet,
    extract_outline_text,
    parse_outline_sheet,
    parse_shot_sheet,
    sheet_plaintext,
    workbook_to_grids_from_path,
)
from services.import_parse_fix import try_fix_parsed_rows
from services.shot_grouping import build_macro_rows, preview_group_stats, suggest_groups

_IMPORT_ROOT = Path("uploads/import_sessions")
_MAX_FILE_BYTES = 20 * 1024 * 1024
_ALLOWED_EXT = {".xlsx", ".xls", ".docx"}

# xls 需 xlrd；v1 仅支持 xlsx/docx
_SUPPORTED_EXT = {".xlsx", ".docx"}


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _ensure_import_root() -> Path:
    _IMPORT_ROOT.mkdir(parents=True, exist_ok=True)
    return _IMPORT_ROOT


def compute_content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def validate_upload_filename(filename: str, size: int) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in _ALLOWED_EXT:
        raise ValueError("仅支持 .xlsx、.xls、.docx 文件")
    if ext == ".xls":
        raise ValueError("暂不支持旧版 .xls，请另存为 .xlsx")
    if ext not in _SUPPORTED_EXT:
        raise ValueError(f"不支持的格式: {ext}")
    if size > _MAX_FILE_BYTES:
        raise ValueError("文件超过 20MB 限制")
    return ext


def save_import_session(file_bytes: bytes, filename: str) -> tuple[str, Path]:
    validate_upload_filename(filename, len(file_bytes))
    session_id = str(uuid.uuid4())
    session_dir = _ensure_import_root() / session_id
    session_dir.mkdir(parents=True, exist_ok=True)
    dest = session_dir / Path(filename).name
    dest.write_bytes(file_bytes)
    meta = {
        "filename": Path(filename).name,
        "ext": dest.suffix.lower(),
        "created_at": _utcnow().isoformat(),
    }
    (session_dir / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return session_id, dest


def get_session_file(session_id: str) -> Path:
    session_dir = _ensure_import_root() / session_id
    if not session_dir.is_dir():
        raise FileNotFoundError("导入会话不存在或已过期")
    meta_path = session_dir / "meta.json"
    if not meta_path.is_file():
        raise FileNotFoundError("导入会话元数据缺失")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    file_path = session_dir / meta["filename"]
    if not file_path.is_file():
        raise FileNotFoundError("导入文件不存在")
    return file_path


def extract_word_text(file_path: Path) -> str:
    doc = Document(str(file_path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n\n".join(parts).strip()


def _load_workbook_grids(file_path: Path) -> dict[str, list[list[str]]]:
    return workbook_to_grids_from_path(str(file_path))


def _get_import_logs(db: Session, project_id: str) -> dict[str, ExcelImportLog]:
    rows = (
        db.query(ExcelImportLog)
        .filter(ExcelImportLog.project_id == project_id)
        .all()
    )
    return {r.sheet_name: r for r in rows}


def scan_document(
    db: Session,
    project_id: str,
    file_bytes: bytes,
    filename: str,
) -> dict[str, Any]:
    session_id, file_path = save_import_session(file_bytes, filename)
    ext = file_path.suffix.lower()
    logs = _get_import_logs(db, project_id)
    sheets: list[dict[str, Any]] = []

    if ext == ".docx":
        text = extract_word_text(file_path)
        content_hash = compute_content_hash(text)
        sheet_name = "__word__"
        prev = logs.get(sheet_name)
        if prev and prev.content_hash == content_hash:
            status = "skipped"
        elif prev:
            status = "changed"
        else:
            status = "new"
        sheets.append(
            {
                "sheet_name": sheet_name,
                "display_name": Path(filename).stem,
                "kind": "outline",
                "status": status,
                "content_hash": content_hash,
                "linked_node_id": prev.linked_node_id if prev else None,
                "preview_chars": len(text),
            }
        )
        return {
            "import_session_id": session_id,
            "file_type": "docx",
            "filename": Path(filename).name,
            "sheets": sheets,
        }

    grids = _load_workbook_grids(file_path)
    for sheet_name, grid in grids.items():
        kind = classify_sheet(grid)
        plaintext = sheet_plaintext(grid)
        content_hash = compute_content_hash(plaintext)
        prev = logs.get(sheet_name)
        if prev and prev.content_hash == content_hash:
            status = "skipped"
        elif prev:
            status = "changed"
        else:
            status = "new"
        preview = ""
        if kind == "outline":
            preview = extract_outline_text(grid)[:200]
        elif kind == "shot_table":
            preview = f"分镜表（约 {len(grid)} 行）"
        sheets.append(
            {
                "sheet_name": sheet_name,
                "display_name": sheet_name,
                "kind": kind,
                "status": status,
                "content_hash": content_hash,
                "linked_node_id": prev.linked_node_id if prev else None,
                "preview": preview,
            }
        )

    return {
        "import_session_id": session_id,
        "file_type": "xlsx",
        "filename": Path(filename).name,
        "sheets": sheets,
    }


def _session_dir(session_id: str) -> Path:
    return _ensure_import_root() / session_id


def _cache_parsed_sheet(session_id: str, sheet_name: str, parsed: dict[str, Any]) -> None:
    safe = sheet_name.replace("/", "_").replace("\\", "_")
    path = _session_dir(session_id) / f"parsed_{safe}.json"
    path.write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")


def load_cached_parsed_sheet(session_id: str, sheet_name: str) -> dict[str, Any] | None:
    safe = sheet_name.replace("/", "_").replace("\\", "_")
    path = _session_dir(session_id) / f"parsed_{safe}.json"
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None


async def parse_document_sheets_async(
    session_id: str,
    sheet_names: list[str],
    *,
    try_llm_fix: bool = True,
) -> list[dict[str, Any]]:
    results = parse_document_sheets(session_id, sheet_names)
    out: list[dict[str, Any]] = []
    for parsed in results:
        if parsed.get("kind") == "shot_table" and not parsed.get("error"):
            attach_self_check(parsed)
            if not parsed.get("self_check", {}).get("ok") and try_llm_fix:
                fix = await try_fix_parsed_rows(parsed, use_llm=True)
                if fix.get("fixed"):
                    parsed["rows"] = fix["rows"]
                    parsed["llm_fix"] = {
                        "fixed": True,
                        "source": fix.get("source"),
                        "fix_summary": fix.get("fix_summary"),
                    }
                    attach_self_check(parsed)
                else:
                    parsed["llm_fix"] = {
                        "fixed": False,
                        "reason": fix.get("reason"),
                    }
            _cache_parsed_sheet(session_id, parsed["sheet_name"], parsed)
        out.append(parsed)
    return out


def group_suggest_for_sheet(
    session_id: str,
    sheet_name: str,
    target_duration: float = 10.0,
) -> dict[str, Any]:
    parsed = load_cached_parsed_sheet(session_id, sheet_name)
    if not parsed:
        raise FileNotFoundError("请先解析该 sheet")
    micro_rows = parsed.get("rows") or []
    groups = suggest_groups(micro_rows, target_duration=target_duration)
    return {
        "sheet_name": sheet_name,
        "groups": groups,
        "preview_stats": preview_group_stats(micro_rows, groups),
        "micro_rows": micro_rows,
        "segments": parsed.get("segments") or [],
        "source": "rule",
        "summary": "",
    }


async def group_suggest_llm_for_sheet(
    session_id: str,
    sheet_name: str,
) -> dict[str, Any]:
    from services.shot_grouping_llm import suggest_groups_llm

    parsed = load_cached_parsed_sheet(session_id, sheet_name)
    if not parsed:
        raise FileNotFoundError("请先解析该 sheet")
    micro_rows = parsed.get("rows") or []
    segments = parsed.get("segments") or []
    result = await suggest_groups_llm(micro_rows, segments)
    groups = result.get("groups") or []
    return {
        "sheet_name": sheet_name,
        "groups": groups,
        "preview_stats": preview_group_stats(micro_rows, groups),
        "micro_rows": micro_rows,
        "segments": segments,
        "source": result.get("source") or "llm",
        "summary": result.get("summary") or "",
    }


def parse_document_sheets(
    session_id: str,
    sheet_names: list[str],
) -> list[dict[str, Any]]:
    file_path = get_session_file(session_id)
    ext = file_path.suffix.lower()
    results: list[dict[str, Any]] = []

    if ext == ".docx":
        text = extract_word_text(file_path)
        results.append(
            {
                "sheet_name": "__word__",
                "kind": "outline",
                "text": text,
                "text_preview": text[:500] + ("…" if len(text) > 500 else ""),
                "stats": {"char_count": len(text)},
                "content_hash": compute_content_hash(text),
            }
        )
        return results

    grids = _load_workbook_grids(file_path)
    for name in sheet_names:
        if name not in grids:
            continue
        grid = grids[name]
        kind = classify_sheet(grid)
        content_hash = compute_content_hash(sheet_plaintext(grid))
        if kind == "shot_table":
            parsed = attach_self_check(parse_shot_sheet(name, grid))
        elif kind == "outline":
            parsed = parse_outline_sheet(name, grid)
        else:
            parsed = {
                "sheet_name": name,
                "kind": "unknown",
                "error": "无法识别 sheet 类型",
            }
        parsed["content_hash"] = content_hash
        results.append(parsed)
    return results


def _parse_canvas_data(project: CanvasProject) -> dict[str, Any]:
    try:
        data = json.loads(project.data or "{}")
        if not isinstance(data, dict):
            return {"nodes": [], "edges": []}
        return {
            "nodes": data.get("nodes") or [],
            "edges": data.get("edges") or [],
        }
    except json.JSONDecodeError:
        return {"nodes": [], "edges": []}


def _find_node_by_sheet(nodes: list[dict], sheet_name: str) -> dict | None:
    for node in nodes:
        meta = (node.get("data") or {}).get("importMeta") or {}
        if meta.get("sheetName") == sheet_name:
            return node
    return None


def _find_outline_node(nodes: list[dict]) -> dict | None:
    for node in nodes:
        if node.get("type") != "text-note":
            continue
        data = node.get("data") or {}
        if data.get("textMode") == "screenplay":
            meta = data.get("importMeta") or {}
            if meta.get("sheetName") == "__word__" or meta.get("kind") == "outline":
                return node
    return None


def _make_node_id(prefix: str) -> str:
    return f"{prefix}-{int(_utcnow().timestamp() * 1000)}-{uuid.uuid4().hex[:6]}"


SCRIPT_TABLE_WIDTH = 1120.0
HORIZONTAL_GAP = 64.0
BASE_X = 120.0
BASE_Y = 160.0


def _node_width(node: dict) -> float:
    if node.get("type") == "script-table":
        return float(node.get("width") or SCRIPT_TABLE_WIDTH)
    return float(node.get("width") or 480.0)


def _outline_position(_nodes: list[dict]) -> dict[str, float]:
    return {"x": BASE_X, "y": BASE_Y}


def _script_table_position(nodes: list[dict], index: int) -> dict[str, float]:
    """分镜表节点横向排列，间距紧凑。"""
    script_tables = [n for n in nodes if n.get("type") == "script-table"]
    if script_tables:
        right_edge = BASE_X
        anchor_y = BASE_Y
        for node in script_tables:
            pos = node.get("position") or {}
            x = float(pos.get("x", BASE_X))
            right_edge = max(right_edge, x + _node_width(node) + HORIZONTAL_GAP)
            anchor_y = float(pos.get("y", BASE_Y))
        return {"x": right_edge + index * (SCRIPT_TABLE_WIDTH + HORIZONTAL_GAP), "y": anchor_y}

    if not nodes:
        return {"x": BASE_X + index * (SCRIPT_TABLE_WIDTH + HORIZONTAL_GAP), "y": BASE_Y}

    right_edge = BASE_X
    anchor_y = BASE_Y
    for node in nodes:
        pos = node.get("position") or {}
        x = float(pos.get("x", BASE_X))
        right_edge = max(right_edge, x + _node_width(node) + HORIZONTAL_GAP)
        anchor_y = float(pos.get("y", BASE_Y))
    return {"x": right_edge + index * (SCRIPT_TABLE_WIDTH + HORIZONTAL_GAP), "y": anchor_y}


def _apply_column_mapping(
    rows: list[dict[str, Any]],
    column_mapping: list[dict[str, Any]] | None,
) -> list[dict[str, Any]]:
    if not column_mapping:
        return rows
    return rows


def _upsert_import_log(
    db: Session,
    project_id: str,
    sheet_name: str,
    content_hash: str,
    linked_node_id: str,
) -> None:
    row = (
        db.query(ExcelImportLog)
        .filter(
            ExcelImportLog.project_id == project_id,
            ExcelImportLog.sheet_name == sheet_name,
        )
        .first()
    )
    now = utcnow()
    if row:
        row.content_hash = content_hash
        row.linked_node_id = linked_node_id
        row.last_imported_at = now
    else:
        db.add(
            ExcelImportLog(
                project_id=project_id,
                sheet_name=sheet_name,
                content_hash=content_hash,
                linked_node_id=linked_node_id,
                last_imported_at=now,
            )
        )


def apply_document_import(
    db: Session,
    project: CanvasProject,
    payload: dict[str, Any],
) -> dict[str, Any]:
    canvas = _parse_canvas_data(project)
    nodes: list[dict] = list(canvas.get("nodes") or [])
    edges: list[dict] = list(canvas.get("edges") or [])
    created_nodes: list[str] = []
    updated_nodes: list[str] = []
    stack_idx = 0

    outline = payload.get("outline")
    if outline and outline.get("confirmed"):
        text = str(outline.get("text") or "").strip()
        sheet_name = str(outline.get("sheet_name") or "__word__")
        content_hash = str(outline.get("content_hash") or compute_content_hash(text))
        replace_id = outline.get("replace_node_id")
        target = _find_outline_node(nodes) if not replace_id else next(
            (n for n in nodes if n.get("id") == replace_id), None
        )
        import_meta = {
            "sheetName": sheet_name,
            "kind": "outline",
            "contentHash": content_hash,
            "importedAt": _utcnow().isoformat(),
        }
        if target:
            target["data"] = {
                **(target.get("data") or {}),
                "content": text,
                "textMode": "screenplay",
                "label": outline.get("label") or target.get("data", {}).get("label") or "剧本",
                "importMeta": import_meta,
            }
            updated_nodes.append(target["id"])
            node_id = target["id"]
        else:
            node_id = _make_node_id("text-note")
            nodes.append(
                {
                    "id": node_id,
                    "type": "text-note",
                    "position": _outline_position(nodes),
                    "zIndex": 1,
                    "data": {
                        "content": text,
                        "textMode": "screenplay",
                        "label": outline.get("label") or "剧本",
                        "zIndex": 1,
                        "importMeta": import_meta,
                    },
                }
            )
            created_nodes.append(node_id)
            stack_idx += 1
        _upsert_import_log(db, project.id, sheet_name, content_hash, node_id)

    shot_tables = payload.get("shot_tables") or []
    if len(shot_tables) > 1:
        raise ValueError("V2 每次仅允许导入一集分镜表")

    for item in shot_tables:
        if not item.get("confirmed"):
            continue
        sheet_name = str(item.get("sheet_name") or "")
        segments = item.get("segments") or []
        content_hash = str(item.get("content_hash") or "")
        replace_id = item.get("replace_node_id")
        groups = item.get("groups")

        rows = item.get("rows") or []
        if groups and not rows:
            parsed = load_cached_parsed_sheet(payload.get("import_session_id") or "", sheet_name)
            micro_rows = (parsed or {}).get("rows") or item.get("micro_rows") or []
            if not micro_rows:
                raise ValueError(f"缺少细分镜数据: {sheet_name}")
            if not segments:
                segments = (parsed or {}).get("segments") or []
            rows = build_macro_rows(micro_rows, groups, segments)
        import_meta = {
            "sheetName": sheet_name,
            "kind": "shot_table",
            "contentHash": content_hash,
            "importedAt": _utcnow().isoformat(),
        }
        existing = None
        if replace_id:
            existing = next((n for n in nodes if n.get("id") == replace_id), None)
        if not existing:
            existing = _find_node_by_sheet(nodes, sheet_name)

        node_data = {
            "label": item.get("label") or sheet_name,
            "rows": rows,
            "segments": segments,
            "globalStyle": "",
            "themeContext": "",
            "continuityMode": True,
            "visualContinuity": False,
            "importMeta": import_meta,
        }

        if existing and existing.get("type") == "script-table":
            existing["data"] = {**(existing.get("data") or {}), **node_data}
            updated_nodes.append(existing["id"])
            node_id = existing["id"]
        else:
            node_id = _make_node_id("script-table")
            nodes.append(
                {
                    "id": node_id,
                    "type": "script-table",
                    "position": _script_table_position(nodes, stack_idx),
                    "zIndex": 1,
                    "width": 1120,
                    "data": node_data,
                }
            )
            created_nodes.append(node_id)
            stack_idx += 1
        _upsert_import_log(db, project.id, sheet_name, content_hash, node_id)

    canvas_data = {"nodes": nodes, "edges": edges}
    project.data = json.dumps(canvas_data, ensure_ascii=False)
    project.updated_at = utcnow()
    db.commit()
    db.refresh(project)

    return {
        "canvas_data": canvas_data,
        "created_node_ids": created_nodes,
        "updated_node_ids": updated_nodes,
    }


def cleanup_session(session_id: str) -> None:
    session_dir = _ensure_import_root() / session_id
    if session_dir.is_dir():
        shutil.rmtree(session_dir, ignore_errors=True)
